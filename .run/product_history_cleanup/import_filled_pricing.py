import json
import math
import time
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/cc/Documents/GitHub/Price")
INPUT_XLSX = Path("/Users/cc/Desktop/历史产品价格_待录入整理.xlsx")
HISTORY_PATH = ROOT / "config" / "quick_pricing_history.json"
PRICING_CONFIG_PATH = ROOT / "config" / "pricing_config.json"

SHIPPING_TIERS = {
    "under4": {"id": "under4", "label": "≤4oz"},
    "4to8": {"id": "4to8", "label": "4–8oz", "fee": 1.77},
    "8to12": {"id": "8to12", "label": "8–12oz", "fee": 2.6},
    "12to16": {"id": "12to16", "label": "12–16oz", "fee": 3.22},
    "1to1_25": {"id": "1to1_25", "label": "1–1.25lb", "fee": 3.72},
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def to_number(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def round_up_cents(value):
    return math.ceil((float(value or 0) - 1e-9) * 100) / 100


def tier_from_fee(value):
    fee = to_number(value)
    if fee is None:
        return None
    if any(abs(fee - item) < 0.01 for item in [0.5, 0.88, 0.99]):
        return SHIPPING_TIERS["under4"]
    if any(abs(fee - item) < 0.01 for item in [1.77, 2.05]):
        return SHIPPING_TIERS["4to8"]
    if any(abs(fee - item) < 0.01 for item in [2.6, 2.84]):
        return SHIPPING_TIERS["8to12"]
    if any(abs(fee - item) < 0.01 for item in [3.22, 3.48]):
        return SHIPPING_TIERS["12to16"]
    if any(abs(fee - item) < 0.01 for item in [3.72, 4.14]):
        return SHIPPING_TIERS["1to1_25"]
    return None


def tier_from_text(value):
    text = clean(value).lower()
    if not text:
        return None
    if "≤4" in text or "<=4" in text or "under4" in text or "4oz" in text:
        return SHIPPING_TIERS["under4"]
    if "4–8" in text or "4-8" in text:
        return SHIPPING_TIERS["4to8"]
    if "8–12" in text or "8-12" in text:
        return SHIPPING_TIERS["8to12"]
    if "12–16" in text or "12-16" in text:
        return SHIPPING_TIERS["12to16"]
    if "1.25" in text:
        return SHIPPING_TIERS["1to1_25"]
    return None


def fee_for_tier(tier, price):
    if not tier:
        return 0
    if tier["id"] == "under4":
        return 0.5 if price <= 3 else 0.88
    return float(tier.get("fee") or 0)


def pricing_factors():
    config = json.loads(PRICING_CONFIG_PATH.read_text(encoding="utf-8"))
    referral = float(config.get("referral_fee_rate", 0.18))
    returns = float(config.get("return_rate", 0.1))
    factor = (1 - returns) * (1 - referral) - returns * referral * referral
    return {
        "currency_rate": float(config.get("currency_rate_rmb_to_usd", 7.2)),
        "first_leg": float(config.get("first_leg_shipping_usd", 0.3)),
        "disposal": float(config.get("disposal_fee_usd", 0.25)),
        "returns": returns,
        "factor": factor,
    }


FACTORS = pricing_factors()


def calculate_price(cost_rmb, tier, margin):
    cost_usd = float(cost_rmb or 0) / FACTORS["currency_rate"]
    price = 2.99
    shipping_fee = fee_for_tier(tier, price)
    for _ in range(6):
        fixed_cost = (
            cost_usd
            + FACTORS["first_leg"]
            + shipping_fee
            + FACTORS["disposal"] * FACTORS["returns"]
        )
        price = fixed_cost / max(0.1, FACTORS["factor"] - margin)
        shipping_fee = fee_for_tier(tier, price)
    fixed_cost = (
        cost_usd
        + FACTORS["first_leg"]
        + shipping_fee
        + FACTORS["disposal"] * FACTORS["returns"]
    )
    raw_price = fixed_cost / max(0.1, FACTORS["factor"] - margin)
    suggested_price = round_up_cents(raw_price)
    return {
        "price": suggested_price,
        "rawPrice": raw_price,
        "shippingFee": fee_for_tier(tier, suggested_price),
        "costUsd": cost_usd,
        "fixedCost": fixed_cost,
    }


def calculate_sku(cost_rmb, tier):
    margin15 = calculate_price(cost_rmb, tier, 0.15)
    margin10 = calculate_price(cost_rmb, tier, 0.10)
    margin5 = calculate_price(cost_rmb, tier, 0.05)
    break_even = calculate_price(cost_rmb, tier, 0)
    return {
        "margin15": margin15["price"],
        "margin10": margin10["price"],
        "margin5": margin5["price"],
        "breakEven": break_even["price"],
        "shippingFee": margin15["shippingFee"],
        "fees": {
            "margin15": margin15["shippingFee"],
            "margin10": margin10["shippingFee"],
            "margin5": margin5["shippingFee"],
            "breakEven": break_even["shippingFee"],
        },
    }


def sheet_rows(workbook, sheet_name):
    sheet = workbook[sheet_name]
    headers = [clean(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def read_import_records():
    workbook = load_workbook(INPUT_XLSX, data_only=True, read_only=True)
    summary_rows = sheet_rows(workbook, "产品汇总")
    detail_rows = sheet_rows(workbook, "SKU明细")
    details_by_id = {}
    for row in detail_rows:
        details_by_id.setdefault(clean(row.get("建议历史ID")), []).append(row)

    now = int(time.time())
    records = []
    for product in summary_rows:
        product_id = clean(product.get("建议历史ID"))
        entered_fee = to_number(product.get("配送费(USD)"))
        tier = tier_from_fee(entered_fee) or tier_from_text(product.get("配送档位"))
        if not tier or entered_fee is None:
            continue

        sku_rows = []
        for detail in details_by_id.get(product_id, []):
            cost_rmb = to_number(detail.get("采购成本(CNY)"))
            if cost_rmb is None:
                continue
            calculated = calculate_sku(cost_rmb, tier)
            sku_rows.append(
                {
                    "sku": clean(detail.get("SKU")),
                    "title": clean(detail.get("品名")) or clean(detail.get("款式/规格")) or "默认款",
                    "costRmb": cost_rmb,
                    "margin15": calculated["margin15"],
                    "margin10": calculated["margin10"],
                    "margin5": calculated["margin5"],
                    "breakEven": calculated["breakEven"],
                    "shippingFee": calculated["shippingFee"],
                    "fees": calculated["fees"],
                    "purchaseLink": clean(detail.get("采购链接")),
                    "linkedQty": to_number(detail.get("关联数量")),
                }
            )

        if not sku_rows:
            continue

        records.append(
            {
                "id": product_id,
                "name": clean(product.get("产品名称")),
                "source": "quick",
                "status": "priced",
                "createdAt": now,
                "updatedAt": now,
                "dateLabel": clean(product.get("日期标签")) or "表格导入",
                "tierId": tier["id"],
                "tierLabel": tier["label"],
                "shippingFee": entered_fee,
                "notes": "由《历史产品价格_待录入整理.xlsx》导入；按当前工作台快速定价规则生成四档利润价格。",
                "rows": sku_rows,
            }
        )
    return records


def fill_existing_record(record):
    tier = tier_from_fee(record.get("shippingFee")) or tier_from_text(record.get("tierLabel"))
    if not tier:
        return False
    changed = False
    for row in record.get("rows", []):
        cost_rmb = to_number(row.get("costRmb"))
        if cost_rmb is None:
            continue
        calculated = calculate_sku(cost_rmb, tier)
        for key in ["margin15", "margin10", "margin5", "breakEven", "shippingFee"]:
            if row.get(key) != calculated[key]:
                row[key] = calculated[key]
                changed = True
        if row.get("fees") != calculated["fees"]:
            row["fees"] = calculated["fees"]
            changed = True
    if changed:
        now = int(time.time())
        record["updatedAt"] = now
        note = "已按当前工作台快速定价规则重算并补齐利润档位。"
        notes = clean(record.get("notes"))
        notes = notes.replace("已按当前工作台快速定价规则补齐缺失利润档位。", "").strip()
        notes = notes.replace(note, "").strip()
        record["notes"] = (notes + " " + note).strip()
    return changed


def main():
    history = json.loads(HISTORY_PATH.read_text(encoding="utf-8"))
    records = history.get("records", [])
    imported = read_import_records()
    imported_ids = {record["id"] for record in imported}

    filled_existing = 0
    retained_records = []
    for record in records:
        if record.get("id") in imported_ids:
            continue
        if fill_existing_record(record):
            filled_existing += 1
        retained_records.append(record)

    history["records"] = retained_records + imported
    HISTORY_PATH.write_text(json.dumps(history, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "importedProducts": len(imported),
                "importedSkuRows": sum(len(record.get("rows", [])) for record in imported),
                "filledExistingRecords": filled_existing,
                "totalRecords": len(history["records"]),
                "importedNames": [record["name"] for record in imported],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
