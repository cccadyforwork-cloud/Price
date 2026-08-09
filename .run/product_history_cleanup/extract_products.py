import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/cc/Documents/GitHub/Price")
SOURCES = [
    {
        "path": Path("/Users/cc/Desktop/采购单和产品管理上传/7月w4采购品上传领星.xlsx"),
        "source_name": "7月W4采购品上传领星",
        "source_type": "new_upload",
        "date_label": "7月W4采购品",
    },
    {
        "path": Path("/Users/cc/Desktop/旧货盘整理/领星同步.xlsx"),
        "source_name": "领星同步",
        "source_type": "legacy_sync",
        "date_label": "领星同步旧货盘",
    },
]


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def number(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def base_and_variant(title):
    title = clean(title)
    if not title:
        return "", ""
    parts = re.split(r"[-－—]", title, maxsplit=1)
    base = parts[0].strip()
    variant = parts[1].strip() if len(parts) > 1 else ""
    return base or title, variant


def offer_id(url):
    match = re.search(r"/offer/(\d+)", clean(url))
    return match.group(1) if match else ""


def read_sheet(source):
    wb = load_workbook(source["path"], data_only=True, read_only=True)
    ws = wb["产品"]
    headers = [clean(cell) for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
        sku = clean(raw.get("*SKU"))
        title = clean(raw.get("品名"))
        if not sku and not title:
            continue
        product_name, variant = base_and_variant(title)
        link = clean(raw.get("采购链接"))
        model = clean(raw.get("型号"))
        offer = offer_id(link)
        if model:
            group_key = f"{source['source_type']}|model|{model}"
            product_code = model
        else:
            product_code = f"LINK-{offer}" if offer else f"LEGACY-{product_name}"
            group_key = f"{source['source_type']}|{product_name}|{offer or link}"
        rows.append(
            {
                "source_name": source["source_name"],
                "source_type": source["source_type"],
                "date_label": source["date_label"],
                "source_file": str(source["path"]),
                "source_sheet": "产品",
                "source_row": excel_row,
                "group_key": group_key,
                "product_code": product_code,
                "product_name": product_name,
                "variant": variant,
                "sku": sku,
                "title": title,
                "linked_qty": number(raw.get("关联数量1")),
                "buyer": clean(raw.get("采购员")),
                "cost_rmb": number(raw.get("采购成本(CNY)")),
                "tax_unit_price": number(raw.get("含税单价")),
                "purchase_link": link,
                "offer_id": offer,
                "supplier": clean(raw.get("供应商名称")),
                "shipping_tier": "",
                "shipping_fee_usd": None,
                "margin15": None,
                "margin10": None,
                "margin5": None,
                "break_even": None,
                "status": "待补配送费",
            }
        )
    return rows


def read_existing_names():
    names = set()
    order_path = ROOT / "config" / "pricing_history_order.json"
    quick_path = ROOT / "config" / "quick_pricing_history.json"
    for path in [order_path, quick_path]:
        if not path.exists():
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        if "products" in data and isinstance(data["products"], dict):
            names.update(data["products"].keys())
        for record in data.get("records", []):
            name = clean(record.get("name"))
            if name:
                names.add(name)
    return sorted(names)


def possible_existing_match(product_name, existing_names):
    matches = []
    for existing in existing_names:
        if not existing or not product_name:
            continue
        if existing in product_name or product_name in existing:
            matches.append(existing)
    return "、".join(matches)


def cost_range(costs):
    values = [cost for cost in costs if cost is not None]
    if not values:
        return ""
    low, high = min(values), max(values)
    if abs(low - high) < 1e-9:
        return round(low, 4)
    return f"{round(low, 4)} - {round(high, 4)}"


def main():
    rows = []
    for source in SOURCES:
        rows.extend(read_sheet(source))

    existing_names = read_existing_names()
    by_group = defaultdict(list)
    for row in rows:
        by_group[row["group_key"]].append(row)

    products = []
    for index, (group_key, group_rows) in enumerate(by_group.items(), start=1):
        first = group_rows[0]
        links = sorted({row["purchase_link"] for row in group_rows if row["purchase_link"]})
        suppliers = sorted({row["supplier"] for row in group_rows if row["supplier"]})
        buyers = sorted({row["buyer"] for row in group_rows if row["buyer"]})
        costs = [row["cost_rmb"] for row in group_rows]
        product_record_id = f"hist-pending-{index:03d}"
        for row in group_rows:
            row["product_record_id"] = product_record_id
            row["existing_match"] = possible_existing_match(row["product_name"], existing_names)
        products.append(
            {
                "product_record_id": product_record_id,
                "source_name": first["source_name"],
                "date_label": first["date_label"],
                "product_code": first["product_code"],
                "product_name": first["product_name"],
                "sku_count": len(group_rows),
                "cost_range_rmb": cost_range(costs),
                "linked_qty_total": sum(row["linked_qty"] or 0 for row in group_rows),
                "purchase_link_count": len(links),
                "purchase_links": "\n".join(links),
                "supplier": "、".join(suppliers),
                "buyer": "、".join(buyers),
                "shipping_tier": "",
                "shipping_fee_usd": None,
                "status": "待补配送费",
                "existing_match": possible_existing_match(first["product_name"], existing_names),
                "notes": "已整理基础资料；配送费和四档价格后续补充。",
            }
        )

    by_overlap = defaultdict(list)
    for row in rows:
        overlap_key = (row["product_name"], row["offer_id"] or row["purchase_link"])
        by_overlap[overlap_key].append(row)

    duplicates = []
    for (product_name, offer), group_rows in by_overlap.items():
        source_names = sorted({row["source_name"] for row in group_rows})
        if len(group_rows) < 2 or len(source_names) < 2:
            continue
        costs = sorted({row["cost_rmb"] for row in group_rows if row["cost_rmb"] is not None})
        duplicates.append(
            {
                "product_name": product_name,
                "offer_id": offer,
                "source_names": "、".join(source_names),
                "sku_count": len(group_rows),
                "skus": "\n".join(row["sku"] for row in group_rows),
                "costs_rmb": "、".join(str(round(cost, 4)) for cost in costs),
                "suggestion": "同一产品/采购链接在两个表均出现；录入前确认是否保留两套SKU，或以新版TTCA SKU为准。"
                if any(row["source_type"] == "new_upload" for row in group_rows)
                else "同一产品重复，录入前核对。",
            }
        )

    summary = {
        "source_files": [str(source["path"]) for source in SOURCES],
        "sku_rows": len(rows),
        "product_records": len(products),
        "duplicate_groups": len(duplicates),
        "existing_history_names": existing_names,
    }
    output = {
        "summary": summary,
        "products": products,
        "sku_rows": rows,
        "duplicates": duplicates,
    }
    out_path = ROOT / ".run" / "product_history_cleanup" / "cleaned_products.json"
    out_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
