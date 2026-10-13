#!/usr/bin/env python3
"""Create an Amazon Price Discounts upload workbook from the official template."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from export_sep_return_refurb_discount import (
    FIXED_DISCOUNT_PRICE,
    load_source,
    prepare_rows,
)


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = ROOT / "input" / "templates" / "amazon_price_discounts.xlsx"
OUTPUT_PATH = ROOT / "output" / "1店价格折扣_9月退货翻新品_新价1.72及25%毛利_可上传.xlsx"
COMMITTED_UNITS = 10

EXPECTED_HEADER_ROW_1 = [
    "SKU",
    "DISCOUNTED PRICE",
    "COMMITTED UNITS",
    "MAX PRICE",
    "MIN PRICE",
    "ERRORS",
    "ERROR DETAILS",
]
EXPECTED_HEADER_ROW_2 = [
    "SKU",
    "折扣价格",
    "确定参与商品数量",
    "最高折扣价格",
    "最低折扣价格",
    "错误",
    "错误详情",
]


def prepare_upload_rows() -> list[tuple[str, float, int]]:
    source_rows, pricing_config, source_tier = load_source()
    fixed_rows, margin_rows = prepare_rows(source_rows, pricing_config, source_tier)
    rows: list[tuple[str, float, int]] = []
    for row in fixed_rows + margin_rows:
        sku = str(row["sku"]).strip()
        if not sku or len(sku) > 40:
            raise ValueError(f"SKU 不符合模板要求：{sku!r}")
        rows.append((sku, row["discount_price"], COMMITTED_UNITS))
    return rows


def build_upload_workbook(rows: list[tuple[str, float, int]]) -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"未找到价格折扣模板：{TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    if workbook.sheetnames != ["说明", "模板", "Data Definitions"]:
        raise ValueError(f"模板工作表结构异常：{workbook.sheetnames}")

    sheet = workbook["模板"]
    header_row_1 = [sheet.cell(1, column).value for column in range(1, 8)]
    header_row_2 = [sheet.cell(2, column).value for column in range(1, 8)]
    if header_row_1 != EXPECTED_HEADER_ROW_1 or header_row_2 != EXPECTED_HEADER_ROW_2:
        raise ValueError("价格折扣模板列标题已变化，停止生成以避免上传失败。")
    if len(rows) > 500:
        raise ValueError(f"SKU 数量超过模板上限：{len(rows)}")

    # 保留官方模板的列标题、说明、格式及预设 500 行，仅替换数据区。
    for row_number in range(3, 503):
        for column in range(1, 8):
            sheet.cell(row_number, column).value = None

    for row_number, (sku, price, committed_units) in enumerate(rows, start=3):
        sheet.cell(row_number, 1).value = sku
        sheet.cell(row_number, 2).value = price
        sheet.cell(row_number, 3).value = committed_units
        sheet.cell(row_number, 2).number_format = "0.00"
        sheet.cell(row_number, 3).number_format = "0"

    workbook.save(OUTPUT_PATH)


if __name__ == "__main__":
    upload_rows = prepare_upload_rows()
    build_upload_workbook(upload_rows)
    fixed_count = sum(1 for _, price, _ in upload_rows if price == FIXED_DISCOUNT_PRICE)
    print(
        f"已生成可上传模板：{len(upload_rows)} 个 SKU，"
        f"其中 $1.72 折扣 {fixed_count} 个、25%毛利价 {len(upload_rows) - fixed_count} 个；"
        f"文件：{OUTPUT_PATH}"
    )
