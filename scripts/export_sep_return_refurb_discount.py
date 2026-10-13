#!/usr/bin/env python3
"""Export September return/refurbished discounts using current pricing."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from recalculate_quick_history import calculate_price, tier_from_record


ROOT = Path(__file__).resolve().parents[1]
HISTORY_PATH = ROOT / "config" / "quick_pricing_history.json"
CONFIG_PATH = ROOT / "config" / "pricing_config.json"
OUTPUT_PATH = ROOT / "output" / "9月退货翻新品_新价1.72及25%毛利折扣表.xlsx"
SOURCE_RECORD_ID = "return-refurb-new-products-20260928-under4"
FIXED_FLOOR_LIMIT = 1.74
FIXED_DISCOUNT_PRICE = 1.72
TARGET_MARGIN = 0.25


def load_source() -> tuple[list[dict], dict, tuple[str, dict]]:
    history = json.loads(HISTORY_PATH.read_text(encoding="utf-8"))
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    record = next(
        (item for item in history["records"] if item.get("id") == SOURCE_RECORD_ID),
        None,
    )
    if record is None:
        raise RuntimeError(f"未找到数据源记录：{SOURCE_RECORD_ID}")
    rows = record.get("rows", [])
    skus = [str(row.get("sku", "")).strip() for row in rows]
    if not skus or any(not sku for sku in skus) or len(set(skus)) != len(skus):
        raise ValueError("来源记录的 SKU 为空或重复，停止生成折扣表。")
    if any(row.get("shippingFeeSource") != "seller-central-backend" for row in rows):
        raise ValueError("来源记录仍有未核对配送费的 SKU，停止生成折扣表。")
    return rows, config, tier_from_record(record)


def profit_at_price_rmb(row: dict, price: float, config: dict) -> float:
    currency_rate = float(config.get("currency_rate_rmb_to_usd", 7.2))
    referral_rate = float(config.get("referral_fee_rate", 0.15))
    return_threshold = float(config.get("return_rate_price_threshold_usd", 3.0))
    configured_return_rate = float(config.get("return_rate", 0.05))
    return_rate = 0.0 if price <= return_threshold else configured_return_rate
    first_leg = float(config.get("first_leg_shipping_usd", 0.4))
    disposal_fee = float(config.get("disposal_fee_usd", 0.0))
    shipping_fee = float(row["shippingFee"])
    price_factor = (1 - return_rate) * (1 - referral_rate) - return_rate * referral_rate * referral_rate
    fixed_cost = (
        float(row.get("costRmb", 0)) / currency_rate
        + first_leg
        + shipping_fee
        + disposal_fee * return_rate
    )
    return (price * price_factor - fixed_cost) * currency_rate


def prepare_rows(source_rows: list[dict], config: dict, tier: tuple[str, dict]) -> tuple[list[dict], list[dict]]:
    fixed_rows: list[dict] = []
    margin_rows: list[dict] = []
    tier_id, tier_info = tier
    for row in source_rows:
        floor_price = float(row["breakEven"])
        recalculated_floor = calculate_price(
            row["costRmb"], tier_id, tier_info, 0, config,
            shipping_fee_override=row["shippingFee"],
        )["price"]
        if abs(recalculated_floor - floor_price) > 0.005:
            raise ValueError(f"{row['sku']} 的保本价与当前规则不一致，停止生成折扣表。")

        if floor_price < FIXED_FLOOR_LIMIT:
            price = FIXED_DISCOUNT_PRICE
            rule = "固定 $1.72 折扣价"
            profit_rmb = profit_at_price_rmb(row, price, config)
            bucket = fixed_rows
        else:
            result = calculate_price(
                row["costRmb"], tier_id, tier_info, TARGET_MARGIN, config,
                shipping_fee_override=row["shippingFee"],
            )
            price = result["price"]
            rule = "25%毛利价"
            profit_rmb = result["profit_rmb"]
            bucket = margin_rows

        if price < floor_price or profit_rmb < -0.005:
            raise ValueError(f"{row['sku']} 的折扣价低于保本价，停止生成折扣表。")

        bucket.append(
            {
                "sku": row.get("sku", ""),
                "title": row.get("title", ""),
                "rule": rule,
                "discount_price": round(price, 2),
                "profit_rmb": round(profit_rmb, 2),
                "floor_price": round(floor_price, 2),
                "shipping_fee": round(float(row["shippingFee"]), 2),
            }
        )

    fixed_rows.sort(key=lambda item: (item["floor_price"], item["sku"]))
    margin_rows.sort(key=lambda item: (item["floor_price"], item["sku"]))
    return fixed_rows, margin_rows


def format_sheet(sheet, title: str, rows: list[dict], subtitle: str) -> None:
    dark_blue = "1F4E78"
    medium_blue = "5B9BD5"
    pale_blue = "D9EAF7"
    pale_red = "FCE4D6"
    dark_red = "9C0006"
    white = "FFFFFF"
    gray = "666666"
    thin_gray = Side(style="thin", color="B7C9D6")

    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:G2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="微软雅黑", size=10, color=gray)
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A3:G3")
    sheet["A3"] = f"合计：{len(rows)} 个 MSKU｜价格与保本价单位：USD｜预计毛利单位：RMB"
    sheet["A3"].font = Font(name="微软雅黑", size=10, color=gray)
    sheet["A3"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["MSKU", "品名", "折扣规则", "折扣价（USD）", "预计毛利（RMB）", "保本价（USD）", "配送费（USD）"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin_gray, bottom=thin_gray, left=thin_gray, right=thin_gray)
    sheet.row_dimensions[4].height = 24

    for row_number, item in enumerate(rows, start=5):
        values = [
            item["sku"],
            item["title"],
            item["rule"],
            item["discount_price"],
            item["profit_rmb"],
            item["floor_price"],
            item["shipping_fee"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(
                horizontal="left" if column in (1, 2, 3) else "center",
                vertical="center",
            )
            cell.border = Border(top=thin_gray, bottom=thin_gray, left=thin_gray, right=thin_gray)
            if row_number % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=pale_blue)

        sheet.cell(row=row_number, column=4).number_format = '$0.00'
        sheet.cell(row=row_number, column=5).number_format = '¥0.00;[Red]-¥0.00'
        sheet.cell(row=row_number, column=6).number_format = '$0.00'
        sheet.cell(row=row_number, column=7).number_format = '$0.00'
        if item["profit_rmb"] < 0:
            profit_cell = sheet.cell(row=row_number, column=5)
            profit_cell.fill = PatternFill("solid", fgColor=pale_red)
            profit_cell.font = Font(name="微软雅黑", size=10, color=dark_red, bold=True)
        sheet.row_dimensions[row_number].height = 22

    for column, width in enumerate([34, 27, 22, 18, 19, 17, 17], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    if rows:
        sheet.auto_filter.ref = f"A4:G{4 + len(rows)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def build_workbook(fixed_rows: list[dict], margin_rows: list[dict]) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "折扣总表"
    all_rows = fixed_rows + margin_rows
    format_sheet(
        overview,
        "9月退货翻新品新价折扣表",
        all_rows,
        "新版规则：保本价低于 $1.74 的定价 $1.72；其余按 25% 毛利价。配送费取已核对的后台费用。",
    )

    fixed_sheet = workbook.create_sheet("1.72固定折扣")
    format_sheet(
        fixed_sheet,
        "9月退货翻新品—$1.72固定折扣",
        fixed_rows,
        "范围：保本价低于 $1.74；折扣售价统一为 $1.72。",
    )

    margin_sheet = workbook.create_sheet("25%毛利折扣")
    format_sheet(
        margin_sheet,
        "9月退货翻新品—25%毛利折扣",
        margin_rows,
        "范围：保本价不低于 $1.74；按当前配送费与 25% 目标毛利重新计算。",
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)


if __name__ == "__main__":
    source_rows, pricing_config, source_tier = load_source()
    fixed_discount_rows, margin_discount_rows = prepare_rows(source_rows, pricing_config, source_tier)
    build_workbook(fixed_discount_rows, margin_discount_rows)
    print(
        f"已导出 {len(fixed_discount_rows) + len(margin_discount_rows)} 行："
        f"$1.72折扣 {len(fixed_discount_rows)} 行，25%毛利折扣 {len(margin_discount_rows)} 行；"
        f"文件：{OUTPUT_PATH}"
    )
