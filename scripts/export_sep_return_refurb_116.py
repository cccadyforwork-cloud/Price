#!/usr/bin/env python3
"""Export September return/refurbished items with floor price below $1.50."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "config" / "quick_pricing_history.json"
OUTPUT = ROOT / "output" / "9月退货翻新品进1.16.xlsx"
SOURCE_RECORD_ID = "return-refurb-new-products-20260928-under4"
FLOOR_PRICE_LIMIT = 1.50


def load_rows() -> list[dict]:
    payload = json.loads(SOURCE.read_text(encoding="utf-8"))
    record = next(
        (item for item in payload["records"] if item.get("id") == SOURCE_RECORD_ID),
        None,
    )
    if record is None:
        raise RuntimeError(f"未找到数据源记录：{SOURCE_RECORD_ID}")

    rows = [
        row
        for row in record.get("rows", [])
        if row.get("breakEven") is not None
        and float(row["breakEven"]) < FLOOR_PRICE_LIMIT
    ]
    rows.sort(key=lambda row: (float(row["breakEven"]), str(row.get("sku", ""))))
    return rows


def build_workbook(rows: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "9月退货翻新品进1.16"
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False

    dark_blue = "1F4E78"
    medium_blue = "5B9BD5"
    pale_blue = "D9EAF7"
    white = "FFFFFF"
    gray = "666666"
    thin_gray = Side(style="thin", color="B7C9D6")

    sheet.merge_cells("A1:E1")
    sheet["A1"] = "9月退货翻新品进1.16"
    sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:E2")
    sheet["A2"] = "筛选条件：9月退货翻新品单中保底价低于 $1.50 的商品"
    sheet["A2"].font = Font(name="微软雅黑", size=10, color=gray)
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A3:E3")
    sheet["A3"] = f"合计：{len(rows)} 个 MSKU｜30%定价与保底价单位：USD｜预计毛利单位：RMB"
    sheet["A3"].font = Font(name="微软雅黑", size=10, color=gray)
    sheet["A3"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["MSKU", "品名", "30%定价（USD）", "预计毛利（RMB）", "保底价（USD）"]
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=col_index, value=header)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin_gray, bottom=thin_gray, left=thin_gray, right=thin_gray)
    sheet.row_dimensions[4].height = 24

    for row_index, item in enumerate(rows, start=5):
        values = [
            item.get("sku", ""),
            item.get("title", ""),
            round(float(item["margin30"]), 2),
            round(float(item["profit30Rmb"]), 2),
            round(float(item["breakEven"]), 2),
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(
                horizontal="left" if col_index in (1, 2) else "center",
                vertical="center",
            )
            cell.border = Border(
                top=thin_gray,
                bottom=thin_gray,
                left=thin_gray,
                right=thin_gray,
            )
            if row_index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=pale_blue)

        sheet.cell(row=row_index, column=3).number_format = '$0.00'
        sheet.cell(row=row_index, column=4).number_format = '¥0.00'
        sheet.cell(row=row_index, column=5).number_format = '$0.00'
        sheet.row_dimensions[row_index].height = 22

    widths = [34, 27, 18, 19, 17]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.auto_filter.ref = f"A4:E{4 + len(rows)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


if __name__ == "__main__":
    selected_rows = load_rows()
    build_workbook(selected_rows)
    print(f"已导出 {len(selected_rows)} 行：{OUTPUT}")
